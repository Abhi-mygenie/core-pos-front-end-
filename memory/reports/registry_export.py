# CR-389: Registry Excel Export — All CRs + BUGs
# Re-runnable. Output: memory/reports/REGISTRY_EXPORT_<YYYY_MM_DD>.xlsx

import json, os, re
from datetime import date
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("Run: pip install openpyxl")

# ── Paths ──────────────────────────────────────────────────────────────────
MEMORY      = Path(__file__).parent.parent
REGISTRY    = MEMORY / "control" / "registry.json"
CR_DIR      = MEMORY / "change_requests"
OUT_FILE    = Path(__file__).parent / f"REGISTRY_EXPORT_{date.today().strftime('%Y_%m_%d')}.xlsx"

# ── Area inference keywords ────────────────────────────────────────────────
AREA_RULES = [
    ("PMS",    ["pms","room","checkin","check-in","checkout","check-out",
                "frontdesk","front-desk","folio","housekeeping","booking",
                "reservation","arrival","departure","laundry","night audit"]),
    ("INV",    ["inventor","stock","purchase","ingredient","recipe",
                "smart purchase","wastage","vendor"]),
    ("CRM",    ["crm","customer","loyalty","wallet","coupon","credit"]),
    ("SHARED", ["sidebar","printer","printing","auth","login","permission",
                "deploy","security","env","settings","notification",
                "socket","firebase","pwa"]),
]

def infer_area(item):
    """Return area string; use registry value if set, else infer from title."""
    if item.get("area"):
        raw = str(item["area"]).upper()
        for label in ("PMS","INV","CRM","SHARED"):
            if label in raw:
                return label
        return "POS"
    title = (str(item.get("title","")) + " " + str(item.get("notes",""))).lower()
    for label, keywords in AREA_RULES:
        if any(k in title for k in keywords):
            return label
    return "POS"

# ── Type normalisation ─────────────────────────────────────────────────────
def norm_type(raw):
    t = str(raw or "").strip().upper()
    if t in ("CR","BUG","GAP","INVESTIGATION"):
        return t
    if t.startswith("CR"):
        return "CR"
    if t.startswith("BUG"):
        return "BUG"
    if t in ("PROD","PROD-HOTFIX","PROD_HOTFIX"):
        return "PROD-HOTFIX"
    return t or "UNKNOWN"

# ── Status normalisation ───────────────────────────────────────────────────
def norm_status(raw):
    u = str(raw or "").upper()
    if "CLOSED" in u:          return "CLOSED"
    if "WONT" in u:            return "WONT-FIX"
    if "DUPLICATE" in u:       return "DUPLICATE"
    if "SUBSUMED" in u or "SUPERSEDED" in u or "RETIRED" in u or "OBSOLETE" in u or "ABSORBED" in u or "SPLIT" in u:
                               return "DUPLICATE"
    if "PARKED" in u or "DEFERRED" in u or "CARRY-FORWARD" in u:
                               return "PARKED"
    if "BLOCKED" in u:         return "BLOCKED"
    # Gate 6 — owner verified / shipped / resolved
    if "GATE_6" in u or "OWNER VERIFIED" in u or "SHIPPED" in u or "RESOLVED" in u or "MAIN VERIFIED" in u:
                               return "CLOSED"
    # Gate 5B — QA verified
    if "GATE_5B" in u or "QA PASS" in u or "QA_PASS" in u or "QA-VERIFIED" in u or "QA VERIFIED" in u:
                               return "QA PASS"
    if "VERIFIED" in u:        return "CLOSED"
    # Gate 5A — code done
    if "GATE_5A" in u or "GATE_5_" in u or "IMPLEMENTED" in u or "FIXED" in u:
                               return "IMPLEMENTED"
    # Gate 4 — coding in progress
    if "GATE_4" in u or "IN PROGRESS" in u:
                               return "GATE 4 GO"
    # Gate 2-3 — planning
    if "GATE_3" in u or "GATE 3" in u or "PLAN_COMPLETE" in u or "GATE_2" in u or "GATE 2" in u or "IMPACT" in u:
                               return "PLANNING"
    # Gate 1 — registered only
    if "INTAKE" in u or "GATE_1" in u or "GATE 1" in u or "REGISTERED" in u or "NOT STARTED" in u:
                               return "INTAKE"
    if "BACKEND" in u:         return "BLOCKED"
    return "OPEN"

def is_closed(status_norm):
    return status_norm in ("CLOSED","WONT-FIX","DUPLICATE","PARKED")

# ── Gate extraction ────────────────────────────────────────────────────────
def extract_gate(item, status_norm):
    g = item.get("gate","")
    if g:
        return str(g)
    mapping = {
        "INTAKE":"1","PLANNING":"2-3","GATE 4 GO":"4",
        "IMPLEMENTED":"5A","QA PASS":"5B",
        "CLOSED":"CLOSED","BLOCKED":"BLOCKED",
        "PARKED":"PARKED","WONT-FIX":"WONT-FIX","DUPLICATE":"DUPLICATE",
    }
    return mapping.get(status_norm, "—")

# ── Description mining ─────────────────────────────────────────────────────
def get_description(item, status_norm):
    """3-5 lines for OPEN, 1 line for CLOSED."""
    iid   = str(item.get("id",""))
    title = str(item.get("title",""))
    notes = str(item.get("notes","") or "")

    if is_closed(status_norm):
        note_snippet = notes.split(".")[0].strip() if notes else ""
        return (title + (f". {note_snippet}" if note_snippet and len(note_snippet) < 80 else ""))[:200]

    intake_path = _find_intake(CR_DIR, iid)
    if intake_path:
        return _mine_intake(intake_path, item)

    base = title
    if notes:
        snippet = notes if len(notes) <= 250 else notes[:250].rsplit(" ",1)[0]
        base = f"{title}. {snippet}"
    return base[:400]

def _find_intake(cr_dir, iid):
    """Return Path to intake doc or None."""
    if not iid or not cr_dir.exists():
        return None
    for f in cr_dir.iterdir():
        name = f.name
        if name.startswith(iid+"_") or name.startswith(iid+"-") or name.startswith(iid+"."):
            return f
    return None

def _mine_intake(path, item):
    """Extract a 3-5 line description from the intake markdown."""
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return str(item.get("title",""))

    lines = []

    m = re.search(r'##\s*(?:\d+\.\s*)?Summary\s*\n(.*?)(?=\n##|\Z)', text, re.S | re.I)
    if m:
        raw = m.group(1).strip()
        paras = [l.lstrip("#-*| ").strip() for l in raw.splitlines() if l.strip() and not l.strip().startswith("|")]
        lines = paras[:5]

    if not lines:
        paras = [l.strip() for l in text.splitlines()
                 if l.strip() and not l.startswith("#") and not l.startswith("|") and len(l.strip()) > 20]
        lines = paras[:3]

    result = " | ".join(lines[:5]) if lines else str(item.get("title",""))
    return result[:500]

# ── Date extraction ────────────────────────────────────────────────────────
def extract_date(item, field_names):
    for f in field_names:
        v = item.get(f,"")
        if v:
            m = re.search(r'\d{4}-\d{2}-\d{2}', str(v))
            return m.group(0) if m else str(v)[:10]
    return "—"

# ── Blocked derivation ─────────────────────────────────────────────────────
def is_blocked(item, status_norm):
    if status_norm == "BLOCKED":
        return "YES"
    if "BLOCKED" in str(item.get("status","")).upper():
        return "YES"
    return "NO"

def blocked_on(item):
    notes = (str(item.get("notes","") or "") + " " + str(item.get("status",""))).lower()
    if "backend" in notes:   return "backend"
    if "credential" in notes or "cred" in notes: return "credentials"
    if "owner" in notes:     return "owner decision"
    if "parent" in notes:    return "parent CR"
    return "—"

# ── Build row ──────────────────────────────────────────────────────────────
COLUMNS = [
    "S.No","ID","Type","Title","Description","Area","Module",
    "Priority","Risk","Gate","Status","Status Detail",
    "Date Registered","Date Implemented","Sprint",
    "Blocked","Blocked On","Files","Related","Notes","Docs",
]

MODULE_MAP = [
    ("Order Entry",         ["order entry","order-entry","order_entry"]),
    ("Menu Management",     ["menu","bulk edit","bulk-edit"]),
    ("Inventory",           ["inventor","stock","ingredient","recipe"]),
    ("Smart Purchase",      ["smart purchase","purchase planner"]),
    ("Reports",             ["report","ledger","settlement","sales","daily"]),
    ("PMS Check-In",        ["check-in","checkin"]),
    ("PMS Bookings",        ["booking","reservation","new booking"]),
    ("PMS Front Desk",      ["front desk","frontdesk"]),
    ("PMS Folio",           ["folio"]),
    ("PMS Housekeeping",    ["housekeeping","laundry"]),
    ("PMS Revenue",         ["revenue","night audit"]),
    ("Printing",            ["print","printer","bill","kot"]),
    ("Sidebar / Nav",       ["sidebar","navigation","nav"]),
    ("Settings",            ["settings","config"]),
    ("Auth / Permissions",  ["auth","login","permission","role"]),
    ("Sockets",             ["socket","realtime","real-time"]),
    ("CRM",                 ["crm","customer","loyalty","wallet"]),
    ("Payments",            ["payment","settle","collect","discount","coupon"]),
]

def _to_str(v):
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return ", ".join(_to_str(x) for x in v)
    if isinstance(v, dict):
        return ", ".join(f"{k}: {_to_str(x)}" for k, x in v.items())
    return str(v)

def build_row(idx, item):
    status_norm = norm_status(item.get("status",""))
    area        = infer_area(item)

    title_lc = str(item.get("title","")).lower()
    module = "—"
    for mod_label, kws in MODULE_MAP:
        if any(k in title_lc for k in kws):
            module = mod_label
            break

    prio = _to_str(item.get("priority") or item.get("severity")) or "—"
    risk = _to_str(item.get("risk")) or "—"

    files_raw = item.get("files", [])
    if isinstance(files_raw, list):
        files_str = ", ".join(_to_str(f) for f in files_raw[:5])
        if len(files_raw) > 5:
            files_str += f" (+{len(files_raw)-5} more)"
    else:
        files_str = _to_str(files_raw)

    related = item.get("related") or item.get("parent") or "—"
    related = _to_str(related) or "—"

    artifacts = item.get("artifact_refs",{})
    if isinstance(artifacts, dict) and artifacts:
        docs = ", ".join(f"{k}: {_to_str(v)}" for k,v in list(artifacts.items())[:3])
    else:
        docs = _to_str(artifacts)[:100] if artifacts else "—"

    return [
        idx,
        _to_str(item.get("id","")),
        norm_type(item.get("type","")),
        _to_str(item.get("title","")),
        get_description(item, status_norm),
        area,
        module,
        prio,
        risk,
        extract_gate(item, status_norm),
        status_norm,
        _to_str(item.get("status",""))[:300],
        extract_date(item, ["registered","created","created_at"]),
        extract_date(item, ["implemented","closed_at","date_implemented"]),
        _to_str(item.get("sprint_key")) or "—",
        is_blocked(item, status_norm),
        blocked_on(item) if is_blocked(item, status_norm)=="YES" else "—",
        files_str or "—",
        related,
        _to_str(item.get("notes",""))[:300],
        docs,
    ]

# ── Style helpers ──────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
ALT_FILL     = PatternFill("solid", fgColor="EBF3FB")

STATUS_COLORS = {
    "CLOSED":       "C6EFCE",
    "QA PASS":      "DDEBF7",
    "IMPLEMENTED":  "D9E1F2",
    "INTAKE":       "FFF2CC",
    "PLANNING":     "FCE4D6",
    "BLOCKED":      "F4CCCC",
    "PARKED":       "E2EFDA",
    "WONT-FIX":     "D9D9D9",
    "DUPLICATE":    "D9D9D9",
}

def style_header(ws):
    for cell in ws[1]:
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def style_rows(ws, status_col_idx):
    for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill_color = ALT_FILL if r_idx % 2 == 0 else None
        status_cell = ws.cell(row=r_idx, column=status_col_idx)
        sc = STATUS_COLORS.get(status_cell.value)
        if sc:
            fill_color = PatternFill("solid", fgColor=sc)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill_color:
                cell.fill = fill_color

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

COL_WIDTHS = [6,14,12,40,60,10,22,10,10,8,14,40,14,14,18,8,14,40,20,40,40]

# ── Summary sheet ──────────────────────────────────────────────────────────
def add_summary(wb, all_rows):
    ws = wb.create_sheet("Summary")
    from collections import Counter

    def section(title, data_rows, col_headers):
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)
        ws.append(col_headers)
        for cell in ws[ws.max_row]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.font = Font(bold=True, color="FFFFFF")
        for row in data_rows:
            ws.append(row)

    combo = Counter((r[2], r[10]) for r in all_rows)
    section("Type × Status", [[k[0], k[1], v] for k,v in sorted(combo.items())],
            ["Type","Status","Count"])

    area_cnt = Counter(r[5] for r in all_rows)
    section("By Area", [[k,v] for k,v in sorted(area_cnt.items(), key=lambda x:-x[1])],
            ["Area","Count"])

    sprint_cnt = Counter(r[14] for r in all_rows)
    section("By Sprint", [[k,v] for k,v in sorted(sprint_cnt.items(), key=lambda x:-x[1])],
            ["Sprint","Count"])

    blocked = [r for r in all_rows if r[15]=="YES"]
    section(f"Blocked Items ({len(blocked)})",
            [[r[1],r[2],r[3][:60],r[16]] for r in blocked],
            ["ID","Type","Title","Blocked On"])

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 20

# ── Main ───────────────────────────────────────────────────────────────────
def main():
    print(f"Reading {REGISTRY} ...")
    with open(REGISTRY, encoding="utf-8") as f:
        data = json.load(f)
    items = data["items"]
    print(f"  {len(items)} items loaded.")

    all_rows = []
    for idx, item in enumerate(items, start=1):
        all_rows.append(build_row(idx, item))
        if idx % 100 == 0:
            print(f"  processed {idx}/{len(items)} ...")

    CLOSED_STATUSES = {"CLOSED","WONT-FIX","DUPLICATE","PARKED"}
    open_rows = [r for r in all_rows if r[10] not in CLOSED_STATUSES]
    print(f"  All: {len(all_rows)} rows | Open: {len(open_rows)} rows")

    wb = openpyxl.Workbook()

    ws_all = wb.active
    ws_all.title = "All Items"
    ws_all.append(COLUMNS)
    for r in all_rows:
        ws_all.append(r)
    ws_all.freeze_panes = "A2"
    ws_all.auto_filter.ref = ws_all.dimensions
    style_header(ws_all)
    status_col = COLUMNS.index("Status") + 1
    style_rows(ws_all, status_col)
    set_col_widths(ws_all, COL_WIDTHS)

    ws_open = wb.create_sheet("Open Only")
    ws_open.append(COLUMNS)
    for r in open_rows:
        ws_open.append(r)
    ws_open.freeze_panes = "A2"
    ws_open.auto_filter.ref = ws_open.dimensions
    style_header(ws_open)
    style_rows(ws_open, status_col)
    set_col_widths(ws_open, COL_WIDTHS)

    add_summary(wb, all_rows)

    wb.save(OUT_FILE)
    print(f"\n✅  Saved: {OUT_FILE}")
    print(f"   All Items: {len(all_rows)} rows")
    print(f"   Open Only: {len(open_rows)} rows")
    print(f"   Summary:   pivots written")

if __name__ == "__main__":
    main()
